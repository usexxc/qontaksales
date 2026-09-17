import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from qontak_sales.apps.accounts.models import Company
from qontak_sales.apps.customers.models import Customer

User = get_user_model()


class CustomerAPITests(APITestCase):
    def setUp(self):
        self.manager = User.objects.create_user(
            username="mgr@a.com", email="mgr@a.com",
            password="***", role="MANAGER",
        )
        self.agent = User.objects.create_user(
            username="agt@a.com", email="agt@a.com",
            password="***", role="AGENT",
        )
        self.other_agent = User.objects.create_user(
            username="agt@b.com", email="agt@b.com",
            password="***", role="AGENT",
        )
        self.cust_own = Customer.objects.create(
            name="Cust A", agent=self.agent
        )
        self.cust_other = Customer.objects.create(
            name="Cust B"
        )

    def test_manager_sees_all(self):
        self.client.force_authenticate(self.manager)
        res = self.client.get("/api/customers/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        names = sorted(c["name"] for c in res.data["results"])
        self.assertEqual(names, ["Cust A", "Cust B"])

    def test_agent_only_sees_assigned(self):
        self.client.force_authenticate(self.agent)
        res = self.client.get("/api/customers/")
        names = [c["name"] for c in res.data["results"]]
        self.assertEqual(names, ["Cust A"])

    def test_agent_cannot_assign_non_agent(self):
        self.client.force_authenticate(self.manager)
        res = self.client.post(
            "/api/customers/",
            {"name": "X", "agent": self.manager.id},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_wilayah_beda_induk_ditolak(self):
        # provinsi Aceh (kode 11) vs kabupaten Bandung (kode 32.xx) -> harus 400
        from qontak_sales.apps.regions.models import Province, Regency

        aceh = Province.objects.create(code="11", name="Aceh")
        jabar = Province.objects.create(code="32", name="Jawa Barat")
        bandung = Regency.objects.create(code="32.04", province=jabar, name="Kabupaten Bandung")

        self.client.force_authenticate(self.manager)
        res = self.client.post(
            "/api/customers/",
            {"name": "X", "province": aceh.id, "regency": bandung.id},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_wilayah_satu_induk_diterima(self):
        from qontak_sales.apps.regions.models import Province, Regency

        aceh = Province.objects.create(code="11", name="Aceh")
        aceh_selatan = Regency.objects.create(code="11.01", province=aceh, name="Aceh Selatan")

        self.client.force_authenticate(self.manager)
        res = self.client.post(
            "/api/customers/",
            {"name": "X", "province": aceh.id, "regency": aceh_selatan.id},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)

    def test_unauthenticated_ditolak(self):
        res = self.client.get("/api/customers/")
        self.assertEqual(res.status_code, status.HTTP_401_UNAUTHORIZED)


class CustomerExportTests(APITestCase):
    def setUp(self):
        self.manager = User.objects.create_user(
            username="mgr@a.com", email="mgr@a.com",
            password="***", role="MANAGER",
        )
        self.agent = User.objects.create_user(
            username="agt@a.com", email="agt@a.com",
            password="***", role="AGENT",
        )
        Customer.objects.create(
            name="Cust A", agent=self.agent, status="CUSTOMER"
        )
        Customer.objects.create(
            name="Cust B", agent=None, status="PROSPECT"
        )
        Customer.objects.create(
            name="Cust C"
        )

    def test_export_xlsx_bisa_dibaca(self):
        self.client.force_authenticate(self.manager)
        res = self.client.get("/api/customers/export/?type=xlsx")
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        wb = load_workbook(io.BytesIO(res.content))
        names = [row[0] for row in wb.active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(names), ["Cust A", "Cust B", "Cust C"])

    def test_export_agent_cuma_milik_sendiri(self):
        self.client.force_authenticate(self.agent)
        res = self.client.get("/api/customers/export/?type=csv")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        content = res.content.decode("utf-8-sig")
        self.assertIn("Cust A", content)
        self.assertNotIn("Cust B", content)
        self.assertNotIn("Cust C", content)

    def test_export_format_ngawur_ditolak(self):
        self.client.force_authenticate(self.manager)
        res = self.client.get("/api/customers/export/?type=pdf")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)


class CustomerImportCsvTests(APITestCase):
    def setUp(self):
        self.manager = User.objects.create_user(
            username="mgr@a.com", email="mgr@a.com",
            password="***", role="MANAGER",
        )

    def _upload(self, content, filename="import.csv", user=None):
        self.client.force_authenticate(user or self.manager)
        return self.client.post(
            "/api/customers/import/",
            {"file": SimpleUploadedFile(filename, content.encode("utf-8-sig"))},
            format="multipart",
        )

    def test_import_csv_normal(self):
        csv_text = (
            "Nama,Nama Perusahaan,Email,Telepon,Alamat,Status\n"
            "Toko Maju,CV Maju,maju@example.com,0811,Jl. A,PROSPECT\n"
            "Toko Sejahtera,CV Sejahtera,sejahtera@example.com,0812,Jl. B,\n"
        )
        res = self._upload(csv_text)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["created"], 2)
        self.assertEqual(res.data["failed"], 0)
        self.assertEqual(Customer.objects.count(), 2)

    def test_import_csv_delimiter_titik_koma(self):
        # Excel di beberapa locale default save CSV pake titik-koma
        csv_text = (
            "Nama;Email;Status\n"
            "Warung Barokah;barokah@example.com;CUSTOMER\n"
        )
        res = self._upload(csv_text)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["created"], 1)
        cust = Customer.objects.get(name="Warung Barokah")
        self.assertEqual(cust.email, "barokah@example.com")
        self.assertEqual(cust.status, "CUSTOMER")

    def test_import_csv_kosong_ditolak(self):
        res = self._upload("")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_import_ekstensi_ngawur_ditolak(self):
        self.client.force_authenticate(self.manager)
        res = self.client.post(
            "/api/customers/import/",
            {"file": SimpleUploadedFile("data.txt", b"Nama\nBudi")},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)


class AuthTests(TestCase):
    def test_weak_password_register_ditolak(self):
        res = self.client.post(
            reverse("register"),
            {
                "name": "T T",
                "email": "t@t.com",
                "password": "***",
                "company_name": "PT T",
            },
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("password", res.data)

    def test_email_duplikat_ditolak(self):
        User.objects.create_user(
            username="dup@x.com", email="dup@x.com", password="***"
        )
        res = self.client.post(
            reverse("register"),
            {
                "name": "D Up",
                "email": "DUP@X.COM",
                "password": "***",
                "company_name": "PT D",
            },
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
